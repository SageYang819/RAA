from io import BytesIO
from typing import Any, Dict, List

from docx import Document
from docx.shared import Pt, Inches


CORE_SECTION_LABELS = {
    "contact",
    "summary",
    "skills",
    "languages",
    "education",
    "experience",
    "projects",
}


def extra_sections_for_docx(final_resume: Dict[str, Any]) -> List[Dict[str, Any]]:
    extra_sections = []
    resume_name = final_resume.get("name", "").strip().lower()

    for section in final_resume.get("sections", []):
        label = section.get("canonical_label", "").strip().lower()
        title = section.get("original_title", "").strip().lower()
        content = [str(item).strip() for item in section.get("content", []) if str(item).strip()]

        if label in CORE_SECTION_LABELS:
            continue

        if title in CORE_SECTION_LABELS:
            continue

        if resume_name:
            if title == resume_name:
                continue
            if len(content) == 1 and content[0].lower() == resume_name:
                continue

        extra_sections.append(section)

    return extra_sections


def set_basic_document_style(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(10.5)


def add_name_and_contact(document: Document, final_resume: Dict[str, Any]) -> None:
    name = final_resume.get("name", "")
    contact_info = final_resume.get("contact_info", [])

    if name:
        p = document.add_paragraph()
        run = p.add_run(name)
        run.bold = True
        run.font.size = Pt(16)

    if contact_info:
        for item in contact_info:
            p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.add_run(item)


def add_heading(document: Document, title: str) -> None:
    document.add_paragraph()
    p = document.add_paragraph()
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(12)


def add_summary(document: Document, final_resume: Dict[str, Any]) -> None:
    add_heading(document, "Summary")
    summary = final_resume.get("summary", "")
    p = document.add_paragraph(summary if summary else "")
    p.paragraph_format.space_after = Pt(3)


def add_skills(document: Document, final_resume: Dict[str, Any]) -> None:
    add_heading(document, "Skills")
    skills = final_resume.get("skills", {})

    if not skills:
        return

    for category, values in skills.items():
        p = document.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        category_run = p.add_run(f"{category.replace('_', ' ').title()}: ")
        category_run.bold = True
        p.add_run(", ".join(values))


def add_languages(document: Document, final_resume: Dict[str, Any]) -> None:
    add_heading(document, "Languages")
    languages = final_resume.get("languages", [])
    p = document.add_paragraph(", ".join(languages) if languages else "")
    p.paragraph_format.space_after = Pt(3)


def add_education(document: Document, final_resume: Dict[str, Any]) -> None:
    add_heading(document, "Education")
    education = final_resume.get("education", [])

    for item in education:
        p = document.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(0)
        p.add_run(item)


def add_experience(document: Document, final_resume: Dict[str, Any]) -> None:
    add_heading(document, "Experience")
    experience = final_resume.get("experience", [])

    for exp in experience:
        header_parts = [exp.get("company", ""), exp.get("location", ""), exp.get("title", "")]
        header = " — ".join([x for x in header_parts if x])
        dates = exp.get("dates", "")
        if dates:
            header = f"{header} ({dates})" if header else f"({dates})"

        if header:
            p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(header)
            run.bold = True

        for bullet in exp.get("bullets", []):
            p = document.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(0)
            p.add_run(bullet)


def add_projects(document: Document, final_resume: Dict[str, Any]) -> None:
    projects = final_resume.get("projects", [])
    if not projects:
        return

    add_heading(document, "Projects")

    for proj in projects:
        name = proj.get("name", "")
        if name:
            p = document.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(name)
            run.bold = True

        for desc in proj.get("description", []):
            p = document.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(0)
            p.add_run(desc)


def add_extra_sections(document: Document, final_resume: Dict[str, Any]) -> None:
    extra_sections = extra_sections_for_docx(final_resume)

    for section in extra_sections:
        title = section.get("original_title") or section.get("canonical_label") or "Other"
        add_heading(document, title)

        for item in section.get("content", []):
            p = document.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(0)
            p.add_run(item)


def build_docx_bytes(final_resume: Dict[str, Any]) -> bytes:
    document = Document()
    set_basic_document_style(document)

    add_name_and_contact(document, final_resume)
    add_summary(document, final_resume)
    add_skills(document, final_resume)
    add_languages(document, final_resume)
    add_education(document, final_resume)
    add_experience(document, final_resume)
    add_projects(document, final_resume)
    add_extra_sections(document, final_resume)

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()