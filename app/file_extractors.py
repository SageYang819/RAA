from io import BytesIO

import pdfplumber
from docx import Document


def extract_text_from_txt(file_bytes: bytes) -> str:
    return file_bytes.decode("utf-8", errors="ignore").strip()


def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = Document(BytesIO(file_bytes))

    parts = []

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)

    for table in doc.tables:
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                cell_text = cell.text.strip()
                if cell_text:
                    row_text.append(cell_text)
            if row_text:
                parts.append(" | ".join(row_text))

    return "\n".join(parts).strip()


def extract_text_from_pdf(file_bytes: bytes) -> str:
    parts = []

    with pdfplumber.open(BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            text = page.extract_text()
            if text:
                parts.append(text)

    return "\n".join(parts).strip()


def extract_resume_text_from_uploaded_file(uploaded_file) -> str:
    if uploaded_file is None:
        return ""

    file_name = uploaded_file.name.lower()
    file_bytes = uploaded_file.read()

    if file_name.endswith(".txt"):
        return extract_text_from_txt(file_bytes)

    if file_name.endswith(".docx"):
        return extract_text_from_docx(file_bytes)

    if file_name.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)

    raise ValueError("Unsupported file type. Please upload a TXT, DOCX, or PDF file.")